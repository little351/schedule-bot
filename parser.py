import glob
import os
import re
from openpyxl import load_workbook

DAY_NAMES = {"Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота", "Воскресенье"}
TIME_RE = re.compile(r"^\d{1,2}\.\d{2}-\d{1,2}\.\d{2}$")  # 8.50-10.20
FIRST_GROUP_COL = 4        # столбец D - первая группа, дальше все столбцы с заголовком
HEADER_MARKER = "День недели"

WEEK_SHEET_RE = re.compile(r"^\d{2}\.\d{2}$")  # листы вида "01.09"


def normalize_group(name) -> str:
    """'13 группа   ' -> '13 группа'"""
    return re.sub(r"\s+", " ", str(name)).strip()


def find_header_row(sheet):
    """Строка с заголовками 'День недели | Дата | Время | группы...' (обычно 7-я)."""
    for row in range(1, 21):
        val = sheet.cell(row=row, column=1).value
        if val and str(val).strip() == HEADER_MARKER:
            return row
    return 7


def _load_workbook(file_path):
    # в некоторых файлах битая ссылка на картинку (xl/drawings/NULL) - openpyxl из-за неё
    # не открывает весь файл, а картинки нам не нужны, поэтому такие ошибки пропускаем
    import openpyxl.reader.excel as xl_reader

    original = xl_reader.find_images

    def safe_find_images(archive, path):
        try:
            return original(archive, path)
        except KeyError:
            return [], []

    xl_reader.find_images = safe_find_images
    try:
        return load_workbook(file_path, data_only=True)
    finally:
        xl_reader.find_images = original


def parse_schedule(file_path: str) -> list[dict]:
    """Возвращает список записей:
    {week_sheet, day, date, time, group, subject}

    День начинается со строки, где в столбце A стоит название дня, и продолжается,
    пока не встретится следующий день. Количество пар в дне не фиксировано: в новых
    файлах у некоторых дней есть 6-я пара (18.30-20.00), и строки после неё сдвигаются.
    """
    wb = _load_workbook(file_path)
    records = []

    for sheet_name in wb.sheetnames:
        if not WEEK_SHEET_RE.match(sheet_name):
            continue  # пропускаем служебные листы

        sheet = wb[sheet_name]
        header_row = find_header_row(sheet)

        # заголовки групп: все непустые столбцы начиная с D
        group_names = {}
        for col in range(FIRST_GROUP_COL, sheet.max_column + 1):
            val = sheet.cell(row=header_row, column=col).value
            if val and str(val).strip():
                group_names[col] = normalize_group(val)

        day_name = None
        date_val = None
        for r in range(header_row + 1, sheet.max_row + 1):
            a_val = sheet.cell(row=r, column=1).value
            if a_val and str(a_val).strip() in DAY_NAMES:
                day_name = str(a_val).strip()
                date_val = sheet.cell(row=r, column=2).value
            if not day_name:
                continue  # строка с направлениями под заголовком

            time_val = sheet.cell(row=r, column=3).value
            if not time_val or not TIME_RE.match(str(time_val).strip()):
                continue

            for col, group in group_names.items():
                cell_val = sheet.cell(row=r, column=col).value
                if cell_val and str(cell_val).strip():
                    records.append({
                        "week_sheet": sheet_name,
                        "day": day_name,
                        "date": date_val,
                        "time": str(time_val).strip(),
                        "group": group,
                        "subject": str(cell_val).strip(),
                    })

    return records


def list_schedule_files(schedule_dir: str) -> list[str]:
    """Все .xlsx в папке (кроме временных файлов Excel '~$...'), по алфавиту."""
    files = glob.glob(os.path.join(schedule_dir, "*.xlsx"))
    return sorted(f for f in files if not os.path.basename(f).startswith("~$"))


def parse_all_schedules(schedule_dir: str) -> list[dict]:
    """Собирает расписание всех групп из всех файлов папки.
    Если одна и та же группа встречается в нескольких файлах, побеждает файл,
    который идёт позже по алфавиту."""
    by_file = [(path, parse_schedule(path)) for path in list_schedule_files(schedule_dir)]

    group_source = {}
    for path, records in by_file:
        for r in records:
            group_source[r["group"]] = path

    result = []
    for path, records in by_file:
        result.extend(r for r in records if group_source[r["group"]] == path)
    return result


if __name__ == "__main__":
    data = parse_all_schedules("data")
    print(f"Всего записей: {len(data)}")
    print("Группы:", sorted({r["group"] for r in data}))
    for r in data[:10]:
        print(r)
