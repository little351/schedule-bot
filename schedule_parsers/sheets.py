"""Единый интерфейс чтения .xlsx и .xls: парсеры работают с Sheet и не знают о формате файла."""
import datetime as dt
import os
import re

SPREADSHEET_EXTENSIONS = (".xlsx", ".xlsm", ".xls")


class Sheet:
    """Лист в виде таблицы значений. Строки и столбцы нумеруются с 1, как в Excel."""

    def __init__(self, name, rows, datemode=0, merged=None):
        self.name = name
        self._rows = rows
        self.datemode = datemode
        self._merged = merged or {}  # (строка, столбец) -> (строка, столбец) левой ячейки объединения
        self.max_row = len(rows)
        self.max_col = max((len(r) for r in rows), default=0)

    def value(self, row, col):
        if 1 <= row <= self.max_row:
            cells = self._rows[row - 1]
            if 1 <= col <= len(cells):
                return cells[col - 1]
        return None

    def text(self, row, col):
        """Значение ячейки как строка без пробелов по краям ('' для пустой)."""
        v = self.value(row, col)
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return re.sub(r"[ \t]+", " ", str(v)).strip()

    def text_merged(self, row, col):
        """Как text(), но ячейка, закрытая горизонтальным объединением, возвращает текст его левой ячейки:
        одна лекция на несколько групп записывается в таблице одной объединённой ячейкой."""
        origin = self._merged.get((row, col))
        if origin and origin[0] == row:
            return self.text(*origin)
        return self.text(row, col)

    def cells(self, max_rows=None):
        """(строка, столбец, текст) для всех непустых ячеек."""
        last = min(self.max_row, max_rows) if max_rows else self.max_row
        for r in range(1, last + 1):
            for c in range(1, len(self._rows[r - 1]) + 1):
                t = self.text(r, c)
                if t:
                    yield r, c, t

    def as_datetime(self, row, col):
        """Дата из ячейки: настоящая дата, число-серийный номер Excel или строка ГГГГ-ММ-ДД."""
        v = self.value(row, col)
        if isinstance(v, dt.datetime):
            return v
        if isinstance(v, dt.date):
            return dt.datetime(v.year, v.month, v.day)
        if isinstance(v, (int, float)) and 20000 < v < 80000:
            epoch = dt.datetime(1904, 1, 1) if self.datemode == 1 else dt.datetime(1899, 12, 30)
            return epoch + dt.timedelta(days=float(v))
        if isinstance(v, str):
            for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d.%m.%Y"):
                try:
                    return dt.datetime.strptime(v.strip(), fmt)
                except ValueError:
                    pass
        return None


class Workbook:
    def __init__(self, path, sheets):
        self.path = path
        self.sheets = sheets  # список Sheet в порядке следования в файле

    @property
    def sheet_names(self):
        return [s.name for s in self.sheets]


def _open_xlsx(path):
    import openpyxl.reader.excel as xl_reader
    from openpyxl import load_workbook

    # в некоторых файлах битая ссылка на картинку (xl/drawings/NULL): openpyxl из-за неё не открывает
    # весь файл, а картинки нам не нужны, поэтому эту ошибку пропускаем
    original = xl_reader.find_images

    def safe_find_images(archive, rel_path):
        try:
            return original(archive, rel_path)
        except KeyError:
            return [], []

    xl_reader.find_images = safe_find_images
    try:
        wb = load_workbook(path, data_only=True)
    finally:
        xl_reader.find_images = original

    sheets = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        merged = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if (r, c) != (rng.min_row, rng.min_col):
                        merged[(r, c)] = (rng.min_row, rng.min_col)
        sheets.append(Sheet(ws.title, rows, merged=merged))
    return Workbook(path, sheets)


def _open_xls(path):
    import xlrd

    book = xlrd.open_workbook(path)
    sheets = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                ctype = sh.cell_type(r, c)
                value = sh.cell_value(r, c)
                if ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    value = None
                elif ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    except Exception:
                        value = None
                elif ctype == xlrd.XL_CELL_TEXT and not value.strip():
                    value = None
                row.append(value)
            rows.append(row)
        sheets.append(Sheet(sh.name, rows, datemode=book.datemode))
    return Workbook(path, sheets)


def open_workbook(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _open_xlsx(path)
    if ext == ".xls":
        return _open_xls(path)
    raise ValueError(f"неподдерживаемый формат файла: {ext or path}")
